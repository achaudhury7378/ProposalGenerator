from agent_framework import ChatAgent, MagenticBuilder
from agent_framework.openai import OpenAIChatClient
from augment.servers.web_searcher import main_researcher
from augment.tools.research_tool import deep_research
import os
import asyncio
import configparser
from datetime import datetime

from docx import Document
from dotenv import load_dotenv

load_dotenv()

_OLLAMA_HOST = os.getenv("OLLAMA_HOST", "http://localhost:11434")
_OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:26b")

config = configparser.RawConfigParser()
config.read('prompts.cfg')


# --- Tools -------------------------------------------------------------------
# Agent Framework consumes plain (async) callables as tools. The schema is
# derived from the type hints and docstring, so no @kernel_function decorator
# or plugin class is needed.

async def web_search(topic: str, prompt: str) -> str:
    """Search and scrape the web for a topic."""
    print(f"\n[TOOL] web_search -> topic={topic!r}")
    results = await main_researcher(topic, prompt)
    print(f"[TOOL] web_search <- {len(results)} tavily/scrape result(s) for {topic!r}")
    for i, r in enumerate(results, 1):
        print(f"  [tavily result {i}] {str(r)[:500]}")
    return "\n\n".join(str(r) for r in results if r)


async def deep_research_fn(topic: str) -> str:
    """Deep multi-query research on a topic."""
    print(f"\n[TOOL] deep_research_fn -> topic={topic!r}")
    report = await deep_research(topic)
    print(f"[TOOL] deep_research_fn <- report ({len(report)} chars) for {topic!r}")
    print(f"  [report preview] {report[:500]}")
    return report


_RESEARCH_TOOLS = [web_search, deep_research_fn]


# --- Chat client -------------------------------------------------------------
# Ollama exposes an OpenAI-compatible API, so we drive it through the
# OpenAIChatClient pointed at the local /v1 endpoint.

def _chat_client() -> OpenAIChatClient:
    return OpenAIChatClient(
        model_id=_OLLAMA_MODEL,
        base_url=f"{_OLLAMA_HOST}/v1",
        api_key="ollama",  # Ollama doesn't validate the key
    )


# --- Worker agents -----------------------------------------------------------

market_researcher = ChatAgent(
    chat_client=_chat_client(),
    name="MarketAnalyst",
    instructions=config['Market Analysis']['prompt'],
    description=config['Market Analysis']['description'],
    tools=_RESEARCH_TOOLS,
)

financial_researcher = ChatAgent(
    chat_client=_chat_client(),
    name="FinancialPlanner",
    instructions=config['Financial Planning']['prompt'],
    description=config['Financial Planning']['description'],
    tools=_RESEARCH_TOOLS,
)

solution_architect = ChatAgent(
    chat_client=_chat_client(),
    name="SolutionDesigner",
    instructions=config['Solution Design']['prompt'],
    description=config['Solution Design']['description'],
    tools=_RESEARCH_TOOLS,
)

customer_researcher = ChatAgent(
    chat_client=_chat_client(),
    name="CustomerEngagement_Analyst",
    instructions=config['Client Engagement']['prompt'],
    description=config['Client Engagement']['description'],
    tools=_RESEARCH_TOOLS,
)

risk_analyst = ChatAgent(
    chat_client=_chat_client(),
    name="RiskManager",
    instructions=config['Risk Management']['prompt'],
    description=config['Risk Management']['description'],
    tools=_RESEARCH_TOOLS,
)


# --- Magentic orchestration --------------------------------------------------
# In Agent Framework the Magentic manager IS the orchestrator: it plans, routes
# work to participants, and synthesizes the final answer. The old "main" planner
# maps to the manager's instructions and the "synthesizer"/Combiner maps to the
# final-answer prompt.

workflow = (
    MagenticBuilder()
    .participants(
        MarketAnalyst=market_researcher,
        FinancialPlanner=financial_researcher,
        SolutionDesigner=solution_architect,
        CustomerEngagement_Analyst=customer_researcher,
        RiskManager=risk_analyst,
    )
    .with_standard_manager(
        chat_client=_chat_client(),
        instructions=config['Main Agent']['prompt'],
        final_answer_prompt=config['Combiner']['prompt'],
        max_round_count=3,
    )
    .build()
)


async def main():
    task = "Create a proposal for supply chain optimization system"
    print(f"[WORKFLOW] starting: {task!r}\n")

    step = 0
    outputs = []
    async for event in workflow.run_stream(task):
        step += 1
        print(f"[STEP {step}] {type(event).__name__}: {str(event)[:800]}")
        data = getattr(event, "data", None)
        if data is not None and type(event).__name__ == "WorkflowOutputEvent":
            outputs.append("".join([texts.text for texts in data.contents]))
            # print(data)

    print("\n[WORKFLOW] complete. Final output(s):\n")
    for output in outputs:
        print(output)

    out_path = _write_docx(task, outputs)
    print(f"\n[WORKFLOW] proposal written to: {out_path}")


def _write_docx(task: str, outputs: list) -> str:
    """Write the accumulated proposal output(s) to a .docx file."""
    doc = Document()
    doc.add_heading("Business Proposal", level=0)
    doc.add_paragraph(f"Task: {task}")
    doc.add_paragraph(f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}")

    if not outputs:
        doc.add_paragraph("No output was produced by the workflow.")

    for output in outputs:
        for line in str(output).splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            else:
                doc.add_paragraph(line)

    filename = f"proposal_{datetime.now():%Y%m%d_%H%M%S}.docx"
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    asyncio.run(main())
