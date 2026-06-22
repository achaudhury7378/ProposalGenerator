from docx import Document
from datetime import datetime
import os
def _write_docx(task: str, outputs: list) -> str:
    """Write the accumulated proposal output(s) to a .docx file."""
    doc = Document()
    doc.add_heading("Business Proposal", level=0)
    doc.add_paragraph(f"Task: {task}")
    doc.add_paragraph(f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}")

    if not outputs:
        doc.add_paragraph("No output was produced by the workflow.")

    for output in outputs:
        for line in str(output).splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            else:
                doc.add_paragraph(line)

    filename = f"proposal_{datetime.now():%Y%m%d_%H%M%S}.docx"
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
    doc.save(out_path)
    return out_path